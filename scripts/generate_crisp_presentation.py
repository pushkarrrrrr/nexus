import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_crisp_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Clean Color Palette
    PRIMARY = RGBColor(15, 23, 42)          # Slate 900
    NAVY = RGBColor(30, 58, 138)            # Blue 900
    ACCENT_BLUE = RGBColor(14, 165, 233)    # Sky 500
    ACCENT_TEAL = RGBColor(13, 148, 136)    # Teal 600
    TEXT_MUTED = RGBColor(100, 116, 139)    # Slate 500
    TEXT_BODY = RGBColor(51, 65, 85)        # Slate 700
    CARD_BG = RGBColor(255, 255, 255)       # White
    CARD_BORDER = RGBColor(226, 232, 240)   # Slate 200
    RED_ACCENT = RGBColor(225, 29, 72)      # Rose 600
    GREEN_ACCENT = RGBColor(16, 185, 129)   # Emerald 500

    LOGO_PATH = "assets/presentation_images/page_1_img_1_X8.png"
    WORKFLOW_IMG = "assets/presentation_images/page_6_img_1_X9.jpg"
    ORCHESTRATION_IMG = "assets/presentation_images/orchestration_diagram.png"
    TECHSTACK_IMG = "assets/presentation_images/page_8_img_1_X10.png"
    ARCH_IMG = "assets/presentation_images/architecture_diagram.png"

    def add_top_bar(slide):
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT_BLUE
        bar.line.fill.background()

    def add_header(slide, title_text, subtitle_text=None):
        add_top_bar(slide)
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.733), Inches(1.0))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Arial"
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = PRIMARY
        
        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.name = "Arial"
            p2.font.size = Pt(12)
            p2.font.color.rgb = TEXT_MUTED
            p2.space_before = Pt(3)

    def add_card(slide, left, top, width, height, title="", border_color=CARD_BORDER, bg_color=CARD_BG):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
        
        if title:
            tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.12), width - Inches(0.4), Inches(0.45))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title
            p.font.name = "Arial"
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = PRIMARY
        return shape

    # ==========================================
    # SLIDE 1: TITLE SLIDE (WITH LOGO)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    add_top_bar(s1)

    # Add College Logo
    if os.path.exists(LOGO_PATH):
        s1.shapes.add_picture(LOGO_PATH, Inches(0.8), Inches(0.4), Inches(1.1), Inches(1.1))

    # College Header Text
    c_box = s1.shapes.add_textbox(Inches(2.1), Inches(0.4), Inches(10.4), Inches(1.1))
    ctf = c_box.text_frame
    ctf.word_wrap = True
    cp1 = ctf.paragraphs[0]
    cp1.text = "RVS COLLEGE OF ENGINEERING AND TECHNOLOGY, JAMSHEDPUR"
    cp1.font.name = "Arial"
    cp1.font.size = Pt(18)
    cp1.font.bold = True
    cp1.font.color.rgb = NAVY

    cp2 = ctf.add_paragraph()
    cp2.text = "DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING"
    cp2.font.name = "Arial"
    cp2.font.size = Pt(13)
    cp2.font.bold = True
    cp2.font.color.rgb = TEXT_MUTED
    cp2.space_before = Pt(2)

    # Main Project Title Box
    p_box = s1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.333), Inches(2.0))
    ptf = p_box.text_frame
    ptf.word_wrap = True
    pp1 = ptf.paragraphs[0]
    pp1.text = "NEXUS:"
    pp1.alignment = PP_ALIGN.CENTER
    pp1.font.name = "Arial"
    pp1.font.size = Pt(44)
    pp1.font.bold = True
    pp1.font.color.rgb = PRIMARY

    pp2 = ptf.add_paragraph()
    pp2.text = "AGENTIC AI OPERATING SYSTEM"
    pp2.alignment = PP_ALIGN.CENTER
    pp2.font.name = "Arial"
    pp2.font.size = Pt(36)
    pp2.font.bold = True
    pp2.font.color.rgb = ACCENT_BLUE
    pp2.space_before = Pt(2)

    pp3 = ptf.add_paragraph()
    pp3.text = "User Intent  ➔  Autonomous Multi-Agent DAGs  ➔  Controlled Computer Execution"
    pp3.alignment = PP_ALIGN.CENTER
    pp3.font.name = "Arial"
    pp3.font.size = Pt(15)
    pp3.font.bold = True
    pp3.font.color.rgb = ACCENT_TEAL
    pp3.space_before = Pt(10)

    # Team & Guide Cards
    add_card(s1, Inches(1.5), Inches(4.7), Inches(5.0), Inches(2.3), "Presented By: Group A5")
    tb_team = s1.shapes.add_textbox(Inches(1.7), Inches(5.2), Inches(4.6), Inches(1.6))
    ttf = tb_team.text_frame
    members = [
        "1. Pushkar Raj (CSE/073/23)",
        "2. Chandan Kr. Sah (CSE/074/23)",
        "3. Prince Tiwari (CSE/075/23)",
        "4. Shruti Kumari (CSE/076/23)"
    ]
    for idx, m in enumerate(members):
        tp = ttf.paragraphs[0] if idx == 0 else ttf.add_paragraph()
        tp.text = m
        tp.font.name = "Arial"
        tp.font.size = Pt(13)
        tp.font.color.rgb = TEXT_BODY
        tp.space_before = Pt(3)

    add_card(s1, Inches(6.8), Inches(4.7), Inches(5.0), Inches(2.3), "Under the Guidance of")
    tb_guide = s1.shapes.add_textbox(Inches(7.0), Inches(5.25), Inches(4.6), Inches(1.5))
    gtf = tb_guide.text_frame
    gp1 = gtf.paragraphs[0]
    gp1.text = "Sulekh Kumar"
    gp1.font.name = "Arial"
    gp1.font.size = Pt(20)
    gp1.font.bold = True
    gp1.font.color.rgb = PRIMARY
    gp2 = gtf.add_paragraph()
    gp2.text = "Assistant Professor\nDept. of Computer Science & Engineering\nRVSCET, Jamshedpur"
    gp2.font.name = "Arial"
    gp2.font.size = Pt(12)
    gp2.font.color.rgb = TEXT_MUTED
    gp2.space_before = Pt(4)

    # ==========================================
    # SLIDE 2: TABLE OF CONTENTS (CRISP & NUMBERED)
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Table of Contents", "Roadmap of today's presentation")

    toc_items = [
        ("01", "Introduction & Problem Statement", "Why chatbots fail and what NEXUS solves"),
        ("02", "System Architecture", "Dual surfaces, FastAPI core & decoupled layers"),
        ("03", "Pictorial Workflow & Safety Pipeline", "6-step execution & human verification loop"),
        ("04", "Multi-Agent Orchestration", "Specialist agents & dynamic self-correction"),
        ("05", "Models Used in NEXUS", "Multi-provider AI Gateway (Gemini, Claude, GPT-4o)"),
        ("06", "Comprehensive Technology Stack", "Next.js 14, Tauri, FastAPI, PostgreSQL, Playwright"),
        ("07", "Policy Engine & 1-Click Rollback", "4-tier risk matrix & SHA-256 snapshot recovery"),
        ("08", "Demerits & System Limitations", "Key engineering constraints & boundary conditions"),
        ("09", "Future Scope & Engineering Roadmap", "Vision perception, cross-platform & hardware keys"),
        ("10", "Conclusion & Key Takeaways", "Capability + Control + Observability + Reversibility")
    ]

    for idx, (num, title, desc) in enumerate(toc_items):
        col = idx // 5
        row = idx % 5
        left_pos = Inches(0.8 + col * 5.95)
        top_pos = Inches(1.6 + row * 1.05)
        
        card = add_card(s2, left_pos, top_pos, Inches(5.6), Inches(0.9))
        tb = s2.shapes.add_textbox(left_pos + Inches(0.2), top_pos + Inches(0.1), Inches(5.2), Inches(0.7))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"{num}. {title}"
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = PRIMARY

        p_sub = tf.add_paragraph()
        p_sub.text = desc
        p_sub.font.name = "Arial"
        p_sub.font.size = Pt(10)
        p_sub.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 3: INTRODUCTION & PROBLEM STATEMENT (CRISP)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Introduction & Problem Statement", "Bridging the gap between AI reasoning and safe real-world execution")

    # Paradigm Banner
    banner = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.6))
    banner.fill.solid()
    banner.fill.fore_color.rgb = RGBColor(238, 242, 255)
    banner.line.color.rgb = RGBColor(199, 210, 254)
    btb = s3.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(11.333), Inches(0.5))
    bp = btb.text_frame.paragraphs[0]
    bp.text = "THE PARADIGM SHIFT: FROM CONVERSATIONAL CHATBOTS  ➔  TO ACTION-ORIENTED AGENTS"
    bp.alignment = PP_ALIGN.CENTER
    bp.font.name = "Arial"
    bp.font.size = Pt(12)
    bp.font.bold = True
    bp.font.color.rgb = NAVY

    # Left: The Problem
    add_card(s3, Inches(0.8), Inches(2.3), Inches(5.6), Inches(4.7), "The Problem: AI Can Think, But Who Controls It?", border_color=RGBColor(254, 205, 211))
    tb_p = s3.shapes.add_textbox(Inches(1.0), Inches(2.9), Inches(5.2), Inches(3.9))
    ptf = tb_p.text_frame
    ptf.word_wrap = True
    prob_items = [
        ("Chatbots Stop at Text", "Generate answers but cannot execute commands, run tests, or modify code."),
        ("Voice Assistants are Shallow", "Siri/Alexa use hardcoded intent templates; zero multi-step planning or reasoning."),
        ("Autonomous Scripts are Hazardous", "Unchecked shell scripts ('bash -c') risk catastrophic file loss and infinite loops."),
        ("No Safety or Undo", "Traditional tools have no pre-execution snapshots or rollback mechanisms.")
    ]
    for idx, (h, b) in enumerate(prob_items):
        p = ptf.paragraphs[0] if idx == 0 else ptf.add_paragraph()
        p.text = f"✖ {h}: "
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = RED_ACCENT
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = b
        run.font.bold = False
        run.font.color.rgb = TEXT_BODY

    # Right: The Objective
    add_card(s3, Inches(6.8), Inches(2.3), Inches(5.6), Inches(4.7), "The Objective: What NEXUS Achieves", border_color=RGBColor(187, 247, 208))
    tb_o = s3.shapes.add_textbox(Inches(7.0), Inches(2.9), Inches(5.2), Inches(3.9))
    otf = tb_o.text_frame
    otf.word_wrap = True
    obj_items = [
        ("Autonomous Planning", "Transforms natural language into structured Directed Acyclic Graphs (DAGs)."),
        ("Air-Gapped Policy Engine", "Every action is evaluated against a 4-Tier Risk Matrix (LOW to CRITICAL)."),
        ("Unified Diff Approvals", "Inspect exact line-by-line file changes before modifying any workspace files."),
        ("1-Click Rollback", "Pre-execution SHA-256 snapshots allow instant recovery if anything goes wrong.")
    ]
    for idx, (h, b) in enumerate(obj_items):
        p = otf.paragraphs[0] if idx == 0 else otf.add_paragraph()
        p.text = f"✔ {h}: "
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = GREEN_ACCENT
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = b
        run.font.bold = False
        run.font.color.rgb = TEXT_BODY

    # ==========================================
    # SLIDE 4: SYSTEM ARCHITECTURE (WITH DIAGRAM)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "System Architecture", "Decoupled multi-tier design separating reasoning, governance, and execution")

    # Insert Architecture Diagram Image
    if os.path.exists(ARCH_IMG):
        s4.shapes.add_picture(ARCH_IMG, Inches(0.8), Inches(1.5), Inches(11.733), Inches(4.5))

    # Crisp Summary Banner below image
    arch_bar = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.15), Inches(11.733), Inches(0.95))
    arch_bar.fill.solid()
    arch_bar.fill.fore_color.rgb = RGBColor(241, 245, 249)
    arch_bar.line.color.rgb = CARD_BORDER
    atb = s4.shapes.add_textbox(Inches(1.0), Inches(6.2), Inches(11.333), Inches(0.8))
    atf = atb.text_frame
    ap = atf.paragraphs[0]
    ap.text = "CORE ARCHITECTURAL HIGHLIGHTS:"
    ap.font.name = "Arial"
    ap.font.size = Pt(11)
    ap.font.bold = True
    ap.font.color.rgb = NAVY

    ap2 = atf.add_paragraph()
    ap2.text = "• Dual Interaction Surfaces: Next.js 14 Dashboard (Web) + Tauri Rust Ambient HUD (Spotlight Overlay)\n• Air-Gapped Trust: Reasoning models have ZERO direct shell access; all operations pass through Policy Engine."
    ap2.font.name = "Arial"
    ap2.font.size = Pt(10)
    ap2.font.color.rgb = TEXT_BODY
    ap2.space_before = Pt(2)

    # ==========================================
    # SLIDE 5: PICTORIAL WORKFLOW (WITH DIAGRAM)
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "AI Agent Safety & Execution Workflow", "The 6-step deterministic pipeline: Snapshot ➔ Policy ➔ Approval ➔ Execution ➔ Verification ➔ Rollback")

    if os.path.exists(WORKFLOW_IMG):
        s5.shapes.add_picture(WORKFLOW_IMG, Inches(0.8), Inches(1.5), Inches(11.733), Inches(4.5))

    # Crisp Explanatory Strip
    wf_bar = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.15), Inches(11.733), Inches(0.95))
    wf_bar.fill.solid()
    wf_bar.fill.fore_color.rgb = RGBColor(241, 245, 249)
    wf_bar.line.color.rgb = CARD_BORDER
    wtb = s5.shapes.add_textbox(Inches(1.0), Inches(6.2), Inches(11.333), Inches(0.8))
    wtf = wtb.text_frame
    wp = wtf.paragraphs[0]
    wp.text = "KEY WORKFLOW INVARIANTS:"
    wp.font.name = "Arial"
    wp.font.size = Pt(11)
    wp.font.bold = True
    wp.font.color.rgb = NAVY

    wp2 = wtf.add_paragraph()
    wp2.text = "1. Pre-execution snapshots capture SHA-256 hashes BEFORE mutation  |  2. Policy Engine halts on HIGH/CRITICAL actions\n3. Git-style unified diffs eliminate ambiguity  |  4. Invariant verification triggers dynamic replanning or 1-click rollback."
    wp2.font.name = "Arial"
    wp2.font.size = Pt(10)
    wp2.font.color.rgb = TEXT_BODY
    wp2.space_before = Pt(2)

    # ==========================================
    # SLIDE 6: MULTI-AGENT ORCHESTRATION (WITH DIAGRAM)
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Multi-Agent Orchestration & Self-Correction", "Specialist agents collaborating via Directed Acyclic Graphs with Phase 15 self-healing")

    if os.path.exists(ORCHESTRATION_IMG):
        s6.shapes.add_picture(ORCHESTRATION_IMG, Inches(0.8), Inches(1.5), Inches(11.733), Inches(4.5))

    # Crisp Summary Strip
    mo_bar = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.15), Inches(11.733), Inches(0.95))
    mo_bar.fill.solid()
    mo_bar.fill.fore_color.rgb = RGBColor(241, 245, 249)
    mo_bar.line.color.rgb = CARD_BORDER
    mtb = s6.shapes.add_textbox(Inches(1.0), Inches(6.2), Inches(11.333), Inches(0.8))
    mtf = mtb.text_frame
    mp = mtf.paragraphs[0]
    mp.text = "AUTONOMOUS SELF-HEALING ENGINE (PHASE 15):"
    mp.font.name = "Arial"
    mp.font.size = Pt(11)
    mp.font.bold = True
    mp.font.color.rgb = NAVY

    mp2 = mtf.add_paragraph()
    mp2.text = "• ErrorClassifier: Identifies runtime errors (TIMEOUT, SYNTAX_ERROR, COMMAND_NOT_FOUND)\n• DAGPatcher: Dynamically splices recovery steps into the active plan with strict replan budgets to avoid loops."
    mp2.font.name = "Arial"
    mp2.font.size = Pt(10)
    mp2.font.color.rgb = TEXT_BODY
    mp2.space_before = Pt(2)

    # ==========================================
    # SLIDE 7: AI MODELS USED IN NEXUS (CRISP & VISUAL)
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "AI Models Used in NEXUS", "Provider-agnostic intelligence unified behind Pydantic v2 structured outputs")

    models_data = [
        ("Google Gemini", "gemini-1.5-pro / 2.0-flash", "Primary Reasoning & Multimodal", [
            "1M+ token context window for full-repo indexing",
            "Native screenshot and document vision analysis",
            "Fastest token throughput for sub-second responses"
        ]),
        ("Anthropic Claude", "claude-3-5-sonnet", "Task Planning & Code Synthesis", [
            "Top-tier DAG plan synthesis & tool accuracy",
            "Deep comprehension of negative constraints",
            "Exceptional Python/TypeScript code generation"
        ]),
        ("OpenAI", "gpt-4o / o1 reasoning", "Mathematical & Algorithmic Logic", [
            "Multi-step algorithmic verification",
            "High reliability in JSON schema conformance",
            "Comparative benchmark evaluation baseline"
        ]),
        ("Local / Ollama", "Llama 3 / Mistral / DeepSeek", "Private & Offline Execution", [
            "Zero cloud telemetry for sensitive corporate data",
            "Runs on local workstation GPU/CPU",
            "Cost-free development and offline fallback"
        ])
    ]

    for i, (m_title, m_ver, m_role, m_points) in enumerate(models_data):
        left_pos = Inches(0.8 + i * 2.98)
        card = add_card(s7, left_pos, Inches(1.6), Inches(2.8), Inches(4.3), m_title)
        tb = s7.shapes.add_textbox(left_pos + Inches(0.15), Inches(2.1), Inches(2.5), Inches(3.6))
        tf = tb.text_frame
        tf.word_wrap = True
        
        vp = tf.paragraphs[0]
        vp.text = m_ver
        vp.font.name = "Arial"
        vp.font.size = Pt(11)
        vp.font.bold = True
        vp.font.color.rgb = ACCENT_BLUE

        rp = tf.add_paragraph()
        rp.text = m_role
        rp.font.name = "Arial"
        rp.font.size = Pt(10)
        rp.font.bold = True
        rp.font.color.rgb = TEXT_MUTED
        rp.space_after = Pt(8)

        for p in m_points:
            bp = tf.add_paragraph()
            bp.text = f"• {p}"
            bp.font.name = "Arial"
            bp.font.size = Pt(10)
            bp.font.color.rgb = TEXT_BODY
            bp.space_before = Pt(5)

    # Bottom Callout: Structured Outputs
    so_bar = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.1), Inches(11.733), Inches(1.0))
    so_bar.fill.solid()
    so_bar.fill.fore_color.rgb = RGBColor(238, 242, 255)
    so_bar.line.color.rgb = RGBColor(199, 210, 254)
    sotb = s7.shapes.add_textbox(Inches(1.0), Inches(6.15), Inches(11.333), Inches(0.9))
    sotf = sotb.text_frame
    sop = sotf.paragraphs[0]
    sop.text = "STRUCTURED OUTPUT & TOKEN LEDGER GUARANTEE:"
    sop.font.name = "Arial"
    sop.font.size = Pt(11)
    sop.font.bold = True
    sop.font.color.rgb = NAVY

    sop2 = sotf.add_paragraph()
    sop2.text = "• Every model response is strictly parsed and validated against Pydantic v2 schemas—guaranteeing 0% JSON syntax failures.\n• Token & Cost Ledger: Real-time calculation of input/output token costs across every execution step."
    sop2.font.name = "Arial"
    sop2.font.size = Pt(10)
    sop2.font.color.rgb = TEXT_BODY
    sop2.space_before = Pt(2)

    # ==========================================
    # SLIDE 8: TECHNOLOGY STACK (WITH LAPTOP DIAGRAM)
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Technology Stack", "Modern full-stack architecture built on a high-performance Python core")

    if os.path.exists(TECHSTACK_IMG):
        s8.shapes.add_picture(TECHSTACK_IMG, Inches(0.8), Inches(1.5), Inches(11.733), Inches(4.5))

    # Crisp Summary Strip
    ts_bar = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.15), Inches(11.733), Inches(0.95))
    ts_bar.fill.solid()
    ts_bar.fill.fore_color.rgb = RGBColor(241, 245, 249)
    ts_bar.line.color.rgb = CARD_BORDER
    tstb = s8.shapes.add_textbox(Inches(1.0), Inches(6.2), Inches(11.333), Inches(0.8))
    tstf = tstb.text_frame
    tsp = tstf.paragraphs[0]
    tsp.text = "FULL-STACK RIGOR & AUTOMATION:"
    tsp.font.name = "Arial"
    tsp.font.size = Pt(11)
    tsp.font.bold = True
    tsp.font.color.rgb = NAVY

    tsp2 = tstf.add_paragraph()
    tsp2.text = "• Frontend: Next.js 14, React 18, TypeScript 5.4, Tailwind CSS  |  Desktop: Tauri 1.5 (Rust core, <100ms launch latency)\n• Backend: FastAPI, SQLAlchemy 2.0, Alembic, PostgreSQL + pgvector  |  Testing: 194 passing Pytest automated test suites."
    tsp2.font.name = "Arial"
    tsp2.font.size = Pt(10)
    tsp2.font.color.rgb = TEXT_BODY
    tsp2.space_before = Pt(2)

    # ==========================================
    # SLIDE 9: POLICY ENGINE & REVERSIBILITY (CRISP)
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Policy Engine & 1-Click Rollback", "Deterministic safety governance and instant state reversion")

    # Left: 4-Tier Risk Matrix
    add_card(s9, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3), "4-Tier Risk Matrix (Air-Gapped Governance)")
    tb9_l = s9.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.5))
    tf9_l = tb9_l.text_frame
    tf9_l.word_wrap = True
    risk_tiers = [
        ("LOW (Autonomous)", "Read-only inspection, file reading, system metrics. Runs automatically without user interruption."),
        ("MEDIUM (Session Grant)", "Non-destructive scripts, running pytest test suites. Authorized via session-scoped permission."),
        ("HIGH (Mandatory HITL)", "Modifying project code, external browser clicks. Requires interactive git-style diff approval."),
        ("CRITICAL (Immediate Confirm)", "Destructive deletes ('rm -rf'), dropping tables. Direct confirmation required; NEVER cached.")
    ]
    for idx, (rt, rd) in enumerate(risk_tiers):
        p = tf9_l.paragraphs[0] if idx == 0 else tf9_l.add_paragraph()
        p.text = f"• {rt}: "
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = RED_ACCENT if "CRITICAL" in rt else (NAVY if "HIGH" in rt else PRIMARY)
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = rd
        run.font.bold = False
        run.font.color.rgb = TEXT_BODY

    # Right: Reversibility Subsystem
    add_card(s9, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.3), "Reversibility & 1-Click Rollback")
    tb9_r = s9.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.2), Inches(4.5))
    tf9_r = tb9_r.text_frame
    tf9_r.word_wrap = True
    rev_steps = [
        ("Pre-Execution SHA-256 Snapshots", "Captures exact byte copies before any tool modifies or deletes a workspace file."),
        ("Live Unified Diff Previews", "Approval modals display exact green (+) additions and red (-) deletions."),
        ("State Drift Detection", "Halts rollback safely if target files were modified by an external process."),
        ("1-Click State Restoration", "Restores original file contents and directory structures cleanly with one click."),
        ("Append-Only Audit Ledger", "Every prompt, decision, tool run, and token cost is permanently saved to an immutable database log.")
    ]
    for idx, (rh, rb) in enumerate(rev_steps):
        p = tf9_r.paragraphs[0] if idx == 0 else tf9_r.add_paragraph()
        p.text = f"✔ {rh}: "
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = GREEN_ACCENT
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = rb
        run.font.bold = False
        run.font.color.rgb = TEXT_BODY

    # ==========================================
    # SLIDE 10: DEMERITS & SYSTEM LIMITATIONS (CRISP)
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Demerits & System Limitations", "Academic transparency: Key trade-offs and engineering constraints")

    demerits_data = [
        ("1. Multi-Step Latency", "Coordination Overhead", [
            "Decomposing tasks into DAGs and policy checks adds 1.5 - 3s latency compared to raw prompt chatbots.",
            "Multi-turn self-reflection loops can increase end-to-end task execution time."
        ]),
        ("2. External Irreversibility", "Network & API Boundaries", [
            "Filesystem mutations are 100% reversible; however, external network calls (sent emails, webhooks) cannot be undone.",
            "Requires strict HIGH/CRITICAL classification for all external mutating calls."
        ]),
        ("3. OS Platform Gating", "Permissions & Accessibility", [
            "Deep desktop control requires explicit macOS Accessibility (AXIsProcessTrusted) and Screen Recording permissions.",
            "Cross-platform parity on Windows & Linux requires containerized subprocess adapters."
        ]),
        ("4. Web DOM Layout Drift", "Browser Automation Fragility", [
            "Playwright browser automation can experience selector failures if third-party web apps deploy sudden UI changes.",
            "Mitigated by combining text matching with accessibility tree resolution."
        ])
    ]

    for i, (title, sub, bullets) in enumerate(demerits_data):
        left_pos = Inches(0.8 + i * 2.98)
        card = add_card(s10, left_pos, Inches(1.6), Inches(2.8), Inches(5.3), title, border_color=RGBColor(254, 205, 211))
        tb = s10.shapes.add_textbox(left_pos + Inches(0.15), Inches(2.2), Inches(2.5), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True
        
        sp = tf.paragraphs[0]
        sp.text = sub
        sp.font.name = "Arial"
        sp.font.size = Pt(11)
        sp.font.bold = True
        sp.font.color.rgb = RED_ACCENT
        sp.space_after = Pt(8)

        for b in bullets:
            bp = tf.add_paragraph()
            bp.text = f"• {b}"
            bp.font.name = "Arial"
            bp.font.size = Pt(10)
            bp.font.color.rgb = TEXT_BODY
            bp.space_before = Pt(6)

    # ==========================================
    # SLIDE 11: FUTURE SCOPE & ROADMAP (CRISP)
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Future Scope & Engineering Roadmap", "Advancing toward ubiquitous, safe agentic computing")

    roadmap_data = [
        ("Phase 16: Academic Benchmark Suite", "Formal Ablation Experiments", [
            "Standardized benchmark test harness across software engineering, sysadmin, and research tasks.",
            "Formal ablation studies: Baseline LLM vs. RAG vs. Memory-RAG vs. Full Agentic NEXUS.",
            "Quantitative evaluation of Task Success Rate, Latency, Cost ($), and 0.0% Safety Violations."
        ]),
        ("Multimodal Vision Screen Perception", "Vision-in-the-Loop Desktop Control", [
            "Expands beyond accessibility trees to real-time multimodal screen vision perception.",
            "Enables navigation of custom HTML Canvas applications, legacy desktop GUIs, and CAD software."
        ]),
        ("Full Cross-Platform Parity", "Linux & Windows Native Subsystems", [
            "Linux integration via Wayland / X11 and native D-Bus system bus.",
            "Windows support via UI Automation APIs and sandboxed PowerShell execution."
        ]),
        ("Hardware-Gated Approvals", "Cryptographic Touch ID / WebAuthn", [
            "Touch ID and WebAuthn hardware security keys to sign CRITICAL-tier actions.",
            "Eliminates stolen-session risks during production cloud deployments and irreversible actions."
        ])
    ]

    for i, (title, sub, bullets) in enumerate(roadmap_data):
        row = 0 if i < 2 else 1
        col = i % 2
        left_pos = Inches(0.8 + col * 5.95)
        top_pos = Inches(1.6 + row * 2.7)
        
        card = add_card(s11, left_pos, top_pos, Inches(5.6), Inches(2.5), title)
        tb = s11.shapes.add_textbox(left_pos + Inches(0.2), top_pos + Inches(0.55), Inches(5.2), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True
        
        sp = tf.paragraphs[0]
        sp.text = sub
        sp.font.name = "Arial"
        sp.font.size = Pt(11)
        sp.font.bold = True
        sp.font.color.rgb = ACCENT_BLUE
        sp.space_after = Pt(4)

        for b in bullets:
            bp = tf.add_paragraph()
            bp.text = f"• {b}"
            bp.font.name = "Arial"
            bp.font.size = Pt(11)
            bp.font.color.rgb = TEXT_BODY
            bp.space_before = Pt(3)

    # ==========================================
    # SLIDE 12: CONCLUSION (CRISP)
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Conclusion & Key Takeaways", "The synthesis of autonomous capability and uncompromised user safety")

    # Main Banner
    concl_banner = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.733), Inches(1.2))
    concl_banner.fill.solid()
    concl_banner.fill.fore_color.rgb = RGBColor(238, 242, 255)
    concl_banner.line.color.rgb = RGBColor(199, 210, 254)
    cbtb = s12.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(0.8))
    cbtf = cbtb.text_frame
    cbp = cbtf.paragraphs[0]
    cbp.text = "“Give AI the ability to act — without giving up control.”"
    cbp.alignment = PP_ALIGN.CENTER
    cbp.font.name = "Arial"
    cbp.font.size = Pt(22)
    cbp.font.bold = True
    cbp.font.color.rgb = NAVY

    cbp2 = cbtf.add_paragraph()
    cbp2.text = "Capability  +  Control  +  Observability  +  Reversibility  =  Trustworthy Agentic AI"
    cbp2.alignment = PP_ALIGN.CENTER
    cbp2.font.name = "Arial"
    cbp2.font.size = Pt(14)
    cbp2.font.bold = True
    cbp2.font.color.rgb = ACCENT_TEAL
    cbp2.space_before = Pt(4)

    concl_cards = [
        ("Proven Engineering Rigor", [
            "15 completed development phases backed by a formal constitution.",
            "194 automated unit and integration tests passing.",
            "Strict Pydantic v2 and TypeScript static type validation."
        ]),
        ("True Agentic Operating System", [
            "Beyond text generation: natively drives terminals, browsers, and desktop apps.",
            "Dynamic DAG replanning and self-correction to recover from failures.",
            "Dual surfaces: Analytical Web Cockpit + Lightweight Ambient HUD."
        ]),
        ("Safe by Architectural Design", [
            "Air-gapped Policy Engine enforcing 4-tier risk classification.",
            "Git-style unified diffs eliminate ambiguity during approvals.",
            "1-click snapshot restoration guarantees protected user workspaces."
        ])
    ]

    for i, (title, bullets) in enumerate(concl_cards):
        left_pos = Inches(0.8 + i * 3.98)
        card = add_card(s12, left_pos, Inches(3.1), Inches(3.7), Inches(3.8), title)
        tb = s12.shapes.add_textbox(left_pos + Inches(0.15), Inches(3.7), Inches(3.4), Inches(3.0))
        tf = tb.text_frame
        tf.word_wrap = True
        
        for idx, b in enumerate(bullets):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = f"✔ {b}"
            p.font.name = "Arial"
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_BODY
            p.space_before = Pt(8)

    # ==========================================
    # SLIDE 13: Q&A AND THANK YOU
    # ==========================================
    s13 = prs.slides.add_slide(blank_layout)
    add_top_bar(s13)

    q_box = s13.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(3.5))
    qtf = q_box.text_frame
    qtf.word_wrap = True
    
    qp1 = qtf.paragraphs[0]
    qp1.text = "Thank You!"
    qp1.alignment = PP_ALIGN.CENTER
    qp1.font.name = "Arial"
    qp1.font.size = Pt(54)
    qp1.font.bold = True
    qp1.font.color.rgb = PRIMARY

    qp2 = qtf.add_paragraph()
    qp2.text = "Questions & Answers"
    qp2.alignment = PP_ALIGN.CENTER
    qp2.font.name = "Arial"
    qp2.font.size = Pt(28)
    qp2.font.bold = True
    qp2.font.color.rgb = ACCENT_BLUE
    qp2.space_before = Pt(8)

    qp3 = qtf.add_paragraph()
    qp3.text = "NEXUS: Agentic AI Operating System\nDepartment of Computer Science & Engineering • RVSCET, Jamshedpur"
    qp3.alignment = PP_ALIGN.CENTER
    qp3.font.name = "Arial"
    qp3.font.size = Pt(14)
    qp3.font.color.rgb = TEXT_MUTED
    qp3.space_before = Pt(14)

    add_card(s13, Inches(2.5), Inches(5.5), Inches(8.333), Inches(1.2), border_color=CARD_BORDER)
    pres_tb = s13.shapes.add_textbox(Inches(2.7), Inches(5.65), Inches(7.933), Inches(0.9))
    ptf_end = pres_tb.text_frame
    ptf_end.word_wrap = True
    
    pp_end = ptf_end.paragraphs[0]
    pp_end.text = "Group A5: Pushkar Raj (073) • Chandan Kr. Sah (074) • Prince Tiwari (075) • Shruti Kumari (076)"
    pp_end.alignment = PP_ALIGN.CENTER
    pp_end.font.name = "Arial"
    pp_end.font.size = Pt(13)
    pp_end.font.bold = True
    pp_end.font.color.rgb = PRIMARY

    pp_end2 = ptf_end.add_paragraph()
    pp_end2.text = "Under the guidance of: Sulekh Kumar (Assistant Professor, CSE)"
    pp_end2.alignment = PP_ALIGN.CENTER
    pp_end2.font.name = "Arial"
    pp_end2.font.size = Pt(12)
    pp_end2.font.color.rgb = TEXT_MUTED
    pp_end2.space_before = Pt(4)

    output_path = "/Users/apple/coding /nexus/NEXUS_Final_Presentation.pptx"
    prs.save(output_path)
    print(f"Crisp presentation successfully saved to {output_path}")

if __name__ == "__main__":
    create_crisp_deck()
