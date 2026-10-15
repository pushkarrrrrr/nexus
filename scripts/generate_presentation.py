import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Professional Color Palette
    BG_COLOR = RGBColor(248, 250, 252)       # Slate 50
    CARD_BG = RGBColor(255, 255, 255)        # Pure White
    CARD_BORDER = RGBColor(226, 232, 240)    # Slate 200
    PRIMARY = RGBColor(15, 23, 42)           # Slate 900
    NAVY = RGBColor(30, 58, 138)             # Blue 900
    ACCENT_BLUE = RGBColor(14, 165, 233)     # Sky 500
    ACCENT_TEAL = RGBColor(13, 148, 136)     # Teal 600
    TEXT_MUTED = RGBColor(100, 116, 139)     # Slate 500
    TEXT_BODY = RGBColor(51, 65, 85)         # Slate 700
    RED_ACCENT = RGBColor(225, 29, 72)       # Rose 600
    GREEN_ACCENT = RGBColor(16, 185, 129)    # Emerald 500

    def add_header(slide, title_text, subtitle_text=None):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Arial"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = PRIMARY
        
        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.name = "Arial"
            p2.font.size = Pt(13)
            p2.font.color.rgb = TEXT_MUTED
            p2.space_before = Pt(4)

    def add_card(slide, left, top, width, height, title="", border_color=CARD_BORDER, bg_color=CARD_BG):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
        
        if title:
            tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), Inches(0.5))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title
            p.font.name = "Arial"
            p.font.size = Pt(15)
            p.font.bold = True
            p.font.color.rgb = PRIMARY
        return shape

    # SLIDE 1: TITLE SLIDE
    s1 = prs.slides.add_slide(blank_layout)
    top_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = ACCENT_BLUE
    top_bar.line.fill.background()

    t_box = s1.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.733), Inches(1.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "RVS COLLEGE OF ENGINEERING AND TECHNOLOGY, JAMSHEDPUR"
    p1.alignment = PP_ALIGN.CENTER
    p1.font.name = "Arial"
    p1.font.size = Pt(20)
    p1.font.bold = True
    p1.font.color.rgb = NAVY

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING"
    p1_sub.alignment = PP_ALIGN.CENTER
    p1_sub.font.name = "Arial"
    p1_sub.font.size = Pt(14)
    p1_sub.font.bold = True
    p1_sub.font.color.rgb = TEXT_MUTED
    p1_sub.space_before = Pt(4)

    proj_box = s1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(2.0))
    ptf = proj_box.text_frame
    ptf.word_wrap = True
    pp1 = ptf.paragraphs[0]
    pp1.text = "NEXUS:"
    pp1.alignment = PP_ALIGN.CENTER
    pp1.font.name = "Arial"
    pp1.font.size = Pt(42)
    pp1.font.bold = True
    pp1.font.color.rgb = PRIMARY

    pp2 = ptf.add_paragraph()
    pp2.text = "AGENTIC AI OPERATING SYSTEM"
    pp2.alignment = PP_ALIGN.CENTER
    pp2.font.name = "Arial"
    pp2.font.size = Pt(36)
    pp2.font.bold = True
    pp2.font.color.rgb = ACCENT_BLUE
    pp2.space_before = Pt(4)

    pp3 = ptf.add_paragraph()
    pp3.text = "Bridging High-Level User Intent to Controlled, Observable & Reversible Computer Execution"
    pp3.alignment = PP_ALIGN.CENTER
    pp3.font.name = "Arial"
    pp3.font.size = Pt(15)
    pp3.font.color.rgb = TEXT_BODY
    pp3.space_before = Pt(10)

    card_team = add_card(s1, Inches(1.5), Inches(4.7), Inches(5.0), Inches(2.2), "Presented By: Group A5")
    team_tb = s1.shapes.add_textbox(Inches(1.7), Inches(5.2), Inches(4.6), Inches(1.6))
    ttf = team_tb.text_frame
    members = [
        "1. Pushkar Raj (CSE/073/23)",
        "2. Chandan Kr. Sah (CSE/074/23)",
        "3. Prince Tiwari (CSE/075/23)",
        "4. Shruti Kumari (CSE/076/23)"
    ]
    for idx, m in enumerate(members):
        tp = ttf.paragraphs[0] if idx == 0 else ttf.add_paragraph()
        tp.text = m
        tp.font.name = "Arial"
        tp.font.size = Pt(13)
        tp.font.color.rgb = TEXT_BODY
        tp.space_before = Pt(3)

    card_guide = add_card(s1, Inches(6.8), Inches(4.7), Inches(5.0), Inches(2.2), "Under the Guidance of")
    guide_tb = s1.shapes.add_textbox(Inches(7.0), Inches(5.3), Inches(4.6), Inches(1.4))
    gtf = guide_tb.text_frame
    gp1 = gtf.paragraphs[0]
    gp1.text = "Sulekh Kumar"
    gp1.font.name = "Arial"
    gp1.font.size = Pt(20)
    gp1.font.bold = True
    gp1.font.color.rgb = PRIMARY
    gp2 = gtf.add_paragraph()
    gp2.text = "Assistant Professor\nDepartment of Computer Science and Engineering\nRVSCET, Jamshedpur"
    gp2.font.name = "Arial"
    gp2.font.size = Pt(12)
    gp2.font.color.rgb = TEXT_MUTED
    gp2.space_before = Pt(4)

    # SLIDE 2: TABLE OF CONTENTS
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Table of Contents", "Comprehensive presentation roadmap across system design, engineering, and evaluation")

    toc_items_left = [
        ("01", "Introduction & Motivation", "The shift from conversational chatbots to agentic operating systems"),
        ("02", "Problem Statement & Objectives", "Bridging the gap between AI reasoning and safe OS execution"),
        ("03", "System Architecture", "Dual-surface interaction, multi-layer decoupling, and trust enclaves"),
        ("04", "Execution Workflow & Safety Pipeline", "Intent analysis, DAG synthesis, air-gapped policy check, and diff approvals"),
        ("05", "Multi-Agent Orchestration", "Specialized agent roles, topological task graphs, and self-correction"),
        ("06", "AI Models & Multi-Provider Gateway", "Gemini, Claude, GPT-4o, Ollama, and Pydantic v2 structured outputs")
    ]
    toc_items_right = [
        ("07", "Comprehensive Technology Stack", "Next.js 14, Tauri/Rust, FastAPI, PostgreSQL/pgvector, Playwright, JXA"),
        ("08", "Policy Engine & Trust Model", "4-tier risk classification, capability catalog, and pre-execution snapshots"),
        ("09", "Reversibility & 1-Click Rollback", "SHA-256 state tracking, unified diff previews, and drift conflict prevention"),
        ("10", "Demerits & System Limitations", "Latency trade-offs, external API boundaries, and OS permission requirements"),
        ("11", "Future Scope & Capstone Roadmap", "Multimodal vision navigation, Linux/Windows parity, and hardware approvals"),
        ("12", "Conclusion & Key Takeaways", "The synthesis of capability, control, observability, and reversibility")
    ]

    for i, (num, title, desc) in enumerate(toc_items_left):
        top_pos = Inches(1.7 + i * 0.88)
        card = add_card(s2, Inches(0.8), top_pos, Inches(5.6), Inches(0.8))
        tb = s2.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.08), Inches(5.2), Inches(0.65))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{num}. {title}"
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = PRIMARY
        p_sub = tf.add_paragraph()
        p_sub.text = desc
        p_sub.font.name = "Arial"
        p_sub.font.size = Pt(10)
        p_sub.font.color.rgb = TEXT_MUTED

    for i, (num, title, desc) in enumerate(toc_items_right):
        top_pos = Inches(1.7 + i * 0.88)
        card = add_card(s2, Inches(6.8), top_pos, Inches(5.6), Inches(0.8))
        tb = s2.shapes.add_textbox(Inches(7.0), top_pos + Inches(0.08), Inches(5.2), Inches(0.65))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{num}. {title}"
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = PRIMARY
        p_sub = tf.add_paragraph()
        p_sub.text = desc
        p_sub.font.name = "Arial"
        p_sub.font.size = Pt(10)
        p_sub.font.color.rgb = TEXT_MUTED

    # SLIDE 3: INTRODUCTION & MOTIVATION
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Introduction & Motivation", "Transitioning from passive text generation to verified computer action")

    banner = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.733), Inches(0.7))
    banner.fill.solid()
    banner.fill.fore_color.rgb = RGBColor(238, 242, 255)
    banner.line.color.rgb = RGBColor(199, 210, 254)
    btb = s3.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(11.333), Inches(0.5))
    btf = btb.text_frame
    bp = btf.paragraphs[0]
    bp.text = "CORE PARADIGM: USER INTENT  ➔  PLANNING (DAG)  ➔  POLICY CHECK  ➔  ACTION  ➔  AUDIT & ROLLBACK"
    bp.alignment = PP_ALIGN.CENTER
    bp.font.name = "Arial"
    bp.font.size = Pt(13)
    bp.font.bold = True
    bp.font.color.rgb = NAVY

    add_card(s3, Inches(0.8), Inches(2.5), Inches(5.6), Inches(4.4), "What is NEXUS?")
    tb_left = s3.shapes.add_textbox(Inches(1.0), Inches(3.1), Inches(5.2), Inches(3.6))
    tf_l = tb_left.text_frame
    tf_l.word_wrap = True
    points_l = [
        ("Agentic AI Operating System", "An autonomous, multi-agent supervisory platform that executes real-world software and system tasks on behalf of the user."),
        ("Not a Chatbot or Wrapper", "Moves beyond web sandboxes to interact natively with terminals, file systems, browsers, and desktop software."),
        ("Strict Operational Constitution", "Guided by AGENTS.md: every action is observable, permission-aware, tamper-audited, and completely reversible."),
        ("Two Unified Interaction Surfaces", "NEXUS Dashboard (Next.js analytical cockpit) and NEXUS Ambient Layer (Tauri lightweight desktop overlay).")
    ]
    for idx, (head, body) in enumerate(points_l):
        p = tf_l.paragraphs[0] if idx == 0 else tf_l.add_paragraph()
        p.text = f"• {head}: "
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = PRIMARY
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = body
        run.font.bold = False
        run.font.color.rgb = TEXT_BODY

    add_card(s3, Inches(6.8), Inches(2.5), Inches(5.6), Inches(4.4), "Why Traditional Approaches Fail")
    tb_right = s3.shapes.add_textbox(Inches(7.0), Inches(3.1), Inches(5.2), Inches(3.6))
    tf_r = tb_right.text_frame
    tf_r.word_wrap = True
    points_r = [
        ("Chatbots Stop at Text", "Models like ChatGPT can suggest code or terminal commands, but require manual copy-pasting, debugging, and verification."),
        ("Voice Assistants are Shallow", "Siri and Google Assistant rely on hardcoded intents (reminders, music). They cannot plan multi-step workflows or reason across tools."),
        ("Unrestricted Agents are Hazardous", "Autonomous CLI scripts executing raw 'bash -c' risk catastrophic file loss, credential leaks, or silent runaway execution loops."),
        ("Lack of Reversibility & Audit", "Existing AI agents provide zero rollback mechanisms. Once a file or database is mutated, recovery is nearly impossible.")
    ]
    for idx, (head, body) in enumerate(points_r):
        p = tf_r.paragraphs[0] if idx == 0 else tf_r.add_paragraph()
        p.text = f"• {head}: "
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = RED_ACCENT if idx == 2 else PRIMARY
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = body
        run.font.bold = False
        run.font.color.rgb = TEXT_BODY

    # SLIDE 4: PROBLEM STATEMENT & OBJECTIVES
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Problem Statement & Project Objectives", "Closing the gap between probabilistic reasoning and deterministic OS control")

    add_card(s4, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1), "The Problem: The Control & Trust Deficit")
    tb4_l = s4.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.3))
    tf4_l = tb4_l.text_frame
    tf4_l.word_wrap = True
    prob_points = [
        ("Execution Disconnect", "Modern foundation models demonstrate human-level reasoning, yet remain isolated from the developer's execution environment."),
        ("Catastrophic Risk Profile", "Giving an LLM direct shell access creates an immediate threat of unintended deletions, syntax bugs, or privilege escalation."),
        ("Opaque Decision Making", "Users cannot inspect what an agent intends to do before it mutates state. Approvals are typically vague ('Allow tool?')."),
        ("Irreversible Side-Effects", "No existing platform takes pre-execution snapshots or offers 1-click state reversion for mutating file/OS operations.")
    ]
    for idx, (head, body) in enumerate(prob_points):
        p = tf4_l.paragraphs[0] if idx == 0 else tf4_l.add_paragraph()
        p.text = f"• {head}: "
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = RED_ACCENT
        p.space_before = Pt(10)
        run = p.add_run()
        run.text = body
        run.font.bold = False
        run.font.color.rgb = TEXT_BODY

    add_card(s4, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.1), "The Objectives: The 4 Tenets of NEXUS")
    tb4_r = s4.shapes.add_textbox(Inches(7.0), Inches(2.4), Inches(5.2), Inches(4.3))
    tf4_r = tb4_r.text_frame
    tf4_r.word_wrap = True
    obj_points = [
        ("Plan & Decompose (DAG Engine)", "Transform open-ended user requests into structured Directed Acyclic Graphs (DAGs) with explicit dependencies and self-correction."),
        ("Govern Every Action (Policy Engine)", "Enforce an air-gapped 4-Tier Risk Matrix (LOW, MEDIUM, HIGH, CRITICAL). No tool execution bypasses policy verification."),
        ("Prove Before Mutating (HITL Diffs)", "Generate git-style unified diffs in real-time approval modals before altering files, codebases, or configuration."),
        ("Guarantee Reversibility (Audit Ledger)", "Capture pre-execution SHA-256 snapshots to enable 1-click state rollback while maintaining an append-only audit trail.")
    ]
    for idx, (head, body) in enumerate(obj_points):
        p = tf4_r.paragraphs[0] if idx == 0 else tf4_r.add_paragraph()
        p.text = f"✔ {head}: "
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = GREEN_ACCENT
        p.space_before = Pt(10)
        run = p.add_run()
        run.text = body
        run.font.bold = False
        run.font.color.rgb = TEXT_BODY

    # SLIDE 5: SYSTEM ARCHITECTURE
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "System Architecture & Security Boundaries", "Strict multi-tier decoupling separating reasoning, governance, execution, and state")

    tiers = [
        ("1. USER SURFACES", "Dual Client Layer", [
            "NEXUS Dashboard (Next.js 14 App Router, WebSockets)",
            "NEXUS Ambient Layer (Tauri + Rust frameless HUD, Cmd+Shift+Space)",
            "Voice Mode & Real-time Anomaly Status Indicators"
        ]),
        ("2. CORE API & WORKER", "FastAPI Asynchronous Gateway", [
            "Async REST API & bidirectional WebSocket event bus",
            "Background Task Worker managing DAG lifecycles",
            "Multi-tenant session isolation and JWT authentication"
        ]),
        ("3. GOVERNANCE & POLICY", "Air-Gapped Trust Engine", [
            "Mandatory air-gap between reasoning agents and OS handles",
            "4-Tier Risk Evaluator (LOW, MEDIUM, HIGH, CRITICAL)",
            "Human-in-the-Loop (HITL) approval lock generation with live diffs"
        ]),
        ("4. EXECUTION & PERSISTENCE", "Controlled OS Tools & Storage", [
            "Sandboxed Subprocess Runner (Directory isolation, buffer caps)",
            "Playwright Persistent Browser & Native macOS JXA / CoreGraphics",
            "PostgreSQL + pgvector / SQLite fallback & 5-Class Memory Taxonomy"
        ])
    ]

    for i, (title, subtitle, bullets) in enumerate(tiers):
        left_pos = Inches(0.8 + i * 2.98)
        card = add_card(s5, left_pos, Inches(1.8), Inches(2.8), Inches(5.1), title)
        tb = s5.shapes.add_textbox(left_pos + Inches(0.15), Inches(2.4), Inches(2.5), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        
        sp = tf.paragraphs[0]
        sp.text = subtitle
        sp.font.name = "Arial"
        sp.font.size = Pt(11)
        sp.font.bold = True
        sp.font.color.rgb = ACCENT_BLUE
        sp.space_after = Pt(8)

        for b in bullets:
            bp = tf.add_paragraph()
            bp.text = f"• {b}"
            bp.font.name = "Arial"
            bp.font.size = Pt(10)
            bp.font.color.rgb = TEXT_BODY
            bp.space_before = Pt(6)

    # SLIDE 6: HOW THE SYSTEM WORKS (WORKFLOW)
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "AI Agent Safety & Execution Workflow", "The 6-step deterministic execution pipeline guaranteeing full accountability")

    steps = [
        ("Step 1", "Snapshot Capture", "Environment & Memory", "Captures pre-execution SHA-256 state of files and resources before any planned modification."),
        ("Step 2", "Policy Check", "Risk Classification", "Air-gapped PolicyEngine validates permissions, checks capability catalog, and evaluates risk tier."),
        ("Step 3", "Human Approval", "Unified Diff Preview", "For HIGH/CRITICAL actions, surfaces interactive diff modal. Blocks execution until approved."),
        ("Step 4", "Controlled Execution", "Sandboxed Runners", "Executes via isolated adapters: directory-jailed terminal, persistent Playwright, or macOS JXA."),
        ("Step 5", "Verify & Self-Heal", "Post-Condition Invariants", "Checks exit codes and outcomes. On failure, ErrorClassifier triggers dynamic DAG replanning."),
        ("Step 6", "Rollback / Audit", "Append-Only Ledger", "Commits immutable event log with token costs. Enables 1-click snapshot restoration if needed.")
    ]

    for i, (s_num, s_title, s_sub, s_desc) in enumerate(steps):
        row = 0 if i < 3 else 1
        col = i % 3
        left_pos = Inches(0.8 + col * 4.0)
        top_pos = Inches(1.8 + row * 2.65)
        
        card = add_card(s6, left_pos, top_pos, Inches(3.7), Inches(2.4), f"{s_num}: {s_title}")
        tb = s6.shapes.add_textbox(left_pos + Inches(0.2), top_pos + Inches(0.55), Inches(3.3), Inches(1.7))
        tf = tb.text_frame
        tf.word_wrap = True
        
        sp = tf.paragraphs[0]
        sp.text = s_sub
        sp.font.name = "Arial"
        sp.font.size = Pt(11)
        sp.font.bold = True
        sp.font.color.rgb = ACCENT_TEAL
        sp.space_after = Pt(4)

        dp = tf.add_paragraph()
        dp.text = s_desc
        dp.font.name = "Arial"
        dp.font.size = Pt(11)
        dp.font.color.rgb = TEXT_BODY
        dp.space_before = Pt(4)

    # SLIDE 7: MULTI-AGENT ORCHESTRATION & SELF-CORRECTION
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Multi-Agent Orchestration & Self-Correction", "Specialized agent collaboration with Phase 15 dynamic error recovery")

    add_card(s7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1), "Specialized Multi-Agent Roles")
    tb7_l = s7.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.3))
    tf7_l = tb7_l.text_frame
    tf7_l.word_wrap = True
    agents = [
        ("OrchestratorAgent", "Supervisor decomposing high-level user intent into structured tasks and delegating work."),
        ("PlanningAgent", "Synthesizes Directed Acyclic Graphs (DAGs) with topological execution ordering."),
        ("ResearchAgent", "Performs hybrid retrieval combining semantic vector similarity with Knowledge Graph traversals."),
        ("DocumentAgent", "Multi-format document ingestion, semantic chunking, and positional text parsing (PDF, MD, CSV, JSON)."),
        ("ComputerAgent", "Interacts with OS tools, executes shell scripts safely, and navigates persistent browser sessions.")
    ]
    for idx, (role, desc) in enumerate(agents):
        p = tf7_l.paragraphs[0] if idx == 0 else tf7_l.add_paragraph()
        p.text = f"• {role}: "
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = PRIMARY
        p.space_before = Pt(6)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_BODY

    add_card(s7, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.1), "Phase 15: Autonomous Self-Healing Engine")
    tb7_r = s7.shapes.add_textbox(Inches(7.0), Inches(2.4), Inches(5.2), Inches(4.3))
    tf7_r = tb7_r.text_frame
    tf7_r.word_wrap = True
    recovery_features = [
        ("Deterministic ErrorClassifier", "Parses runtime exceptions into structured categories (TIMEOUT, SYNTAX_ERROR, COMMAND_NOT_FOUND, PERMISSION_DENIED)."),
        ("SelfCorrectionReflector", "Diagnoses root failure causes using structured LLM reasoning and generates targeted recovery proposals."),
        ("DAGPatcher Subsystem", "Dynamically splices recovery steps into the active task graph while preserving upstream and downstream dependencies."),
        ("Bounded Replan Budgets", "Strictly caps consecutive recovery attempts to eliminate runaway execution loops and excessive API costs."),
        ("Automatic State Rollback", "If recovery cannot be achieved, rolls back modified workspace files to the pre-execution snapshot automatically.")
    ]
    for idx, (feat, desc) in enumerate(recovery_features):
        p = tf7_r.paragraphs[0] if idx == 0 else tf7_r.add_paragraph()
        p.text = f"✔ {feat}: "
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE
        p.space_before = Pt(6)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_BODY

    # SLIDE 8: AI MODELS & MULTI-PROVIDER GATEWAY
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "AI Models Used & The Multi-Provider Gateway", "Provider-agnostic intelligence with strict Pydantic v2 schema enforcement")

    models = [
        ("Google Gemini", "gemini-1.5-pro / 2.0-flash", "Primary Reasoning & Multimodal", [
            "Massive context window (up to 1M tokens)",
            "Native multimodal ingestion of screenshots and documents",
            "High token throughput with cost-efficient execution"
        ]),
        ("Anthropic Claude", "claude-3-5-sonnet", "Complex Planning & Code Synthesis", [
            "Industry-leading task decomposition and DAG synthesis",
            "High reliability in complex tool-use workflows",
            "Exceptional adherence to negative system constraints"
        ]),
        ("OpenAI", "gpt-4o / o1 reasoning", "Mathematical & Deterministic Logic", [
            "Advanced multi-step algorithmic validation",
            "Strict JSON-schema mode conformance",
            "Benchmarking and comparative ablation evaluations"
        ]),
        ("Local / Ollama", "Llama 3 / Mistral / DeepSeek", "Private & Offline Execution", [
            "Zero cloud telemetry for sensitive enterprise data",
            "Runs on local workstation hardware via Ollama REST API",
            "Cost-free local development and offline fallback"
        ])
    ]

    for i, (m_title, m_ver, m_role, m_points) in enumerate(models):
        left_pos = Inches(0.8 + i * 2.98)
        card = add_card(s8, left_pos, Inches(1.8), Inches(2.8), Inches(5.1), m_title)
        tb = s8.shapes.add_textbox(left_pos + Inches(0.15), Inches(2.4), Inches(2.5), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        
        vp = tf.paragraphs[0]
        vp.text = m_ver
        vp.font.name = "Arial"
        vp.font.size = Pt(11)
        vp.font.bold = True
        vp.font.color.rgb = ACCENT_BLUE
        
        rp = tf.add_paragraph()
        rp.text = m_role
        rp.font.name = "Arial"
        rp.font.size = Pt(10)
        rp.font.bold = True
        rp.font.color.rgb = TEXT_MUTED
        rp.space_after = Pt(8)

        for p in m_points:
            bp = tf.add_paragraph()
            bp.text = f"• {p}"
            bp.font.name = "Arial"
            bp.font.size = Pt(10)
            bp.font.color.rgb = TEXT_BODY
            bp.space_before = Pt(5)

    # SLIDE 9: COMPREHENSIVE TECHNOLOGY STACK
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Comprehensive Technology Stack", "Modern, scalable full-stack ecosystem centered around a Python core")

    tech_categories = [
        ("User Surfaces", [
            ("Next.js 14", "App Router web dashboard with server rendering"),
            ("React 18 & TS 5.4", "Component architecture with strict static typing"),
            ("Tauri 1.5 (Rust)", "Native desktop wrapper with global hotkeys (<100ms)"),
            ("Tailwind CSS", "Utility-first modern dark-mode & glassmorphism"),
            ("Lucide Icons", "Consistent SVG visual iconography")
        ]),
        ("Backend & Engine", [
            ("Python 3.11+", "Asynchronous multi-agent execution pipeline"),
            ("FastAPI", "High-performance async REST & WebSocket routing"),
            ("Pydantic v2", "Strict runtime data validation & schema contracts"),
            ("SQLAlchemy 2.0", "Asynchronous ORM with connection pooling"),
            ("Alembic", "Dual-dialect database schema migrations")
        ]),
        ("Data & Memory", [
            ("PostgreSQL", "Relational persistence with ACID guarantees"),
            ("pgvector", "High-dimensional vector embeddings for RAG"),
            ("SQLite (WAL)", "Zero-config local fallback with cosine dot product"),
            ("Redis", "Background task queues & distributed pub/sub"),
            ("Knowledge Graph", "Recursive CTE relational traversal")
        ]),
        ("OS & Tools", [
            ("Playwright", "Persistent Chromium automation behind SSRF guards"),
            ("macOS JXA", "AppleScript/JXA UI element traversal"),
            ("Quartz / CoreGraphics", "Native ctypes mouse & keyboard event synthesis"),
            ("Pytest", "194 automated unit and integration tests passing"),
            ("Ruff & Mypy", "Fast linting, formatting, and strict type checking")
        ])
    ]

    for i, (cat_title, tech_list) in enumerate(tech_categories):
        left_pos = Inches(0.8 + i * 2.98)
        card = add_card(s9, left_pos, Inches(1.8), Inches(2.8), Inches(5.1), cat_title)
        tb = s9.shapes.add_textbox(left_pos + Inches(0.15), Inches(2.4), Inches(2.5), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        
        for idx, (t_name, t_desc) in enumerate(tech_list):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = f"{t_name}: "
            p.font.name = "Arial"
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = PRIMARY
            p.space_before = Pt(6)
            run = p.add_run()
            run.text = t_desc
            run.font.bold = False
            run.font.size = Pt(10)
            run.font.color.rgb = TEXT_BODY

    # SLIDE 10: POLICY ENGINE & REVERSIBILITY
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Policy Engine, Trust Model & Reversibility", "Governing agent actions through explicit risk tiers and 1-click state reversion")

    add_card(s10, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1), "4-Tier Risk Matrix (Air-Gapped Governance)")
    tb10_l = s10.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.3))
    tf10_l = tb10_l.text_frame
    tf10_l.word_wrap = True
    risks = [
        ("LOW Risk (Autonomous)", "Read-only operations, system metrics, file inspection, knowledge search. Executed autonomously without user interruption."),
        ("MEDIUM Risk (Session Grant)", "Non-destructive actions, running local unit tests (pytest), temporary file writes. Authorized via session-scoped permission."),
        ("HIGH Risk (Mandatory HITL)", "Modifying project code, external browser clicks, native desktop control. Requires explicit user approval with git-style unified diff."),
        ("CRITICAL Risk (Immediate Confirm)", "Destructive deletions (rm -rf), dropping databases, modifying security keys. Direct immediate confirmation required; never cached.")
    ]
    for idx, (r_title, r_desc) in enumerate(risks):
        p = tf10_l.paragraphs[0] if idx == 0 else tf10_l.add_paragraph()
        p.text = f"• {r_title}: "
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = RED_ACCENT if "CRITICAL" in r_title else (NAVY if "HIGH" in r_title else PRIMARY)
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = r_desc
        run.font.bold = False
        run.font.color.rgb = TEXT_BODY

    add_card(s10, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.1), "Side-Effect Journaling & 1-Click Rollback")
    tb10_r = s10.shapes.add_textbox(Inches(7.0), Inches(2.4), Inches(5.2), Inches(4.3))
    tf10_r = tb10_r.text_frame
    tf10_r.word_wrap = True
    rev_points = [
        ("Pre-Execution SHA-256 Snapshots", "Before any tool modifies or deletes a file, an exact byte snapshot and checksum are recorded in the local snapshot store."),
        ("Git-Style Unified Diff Previews", "Approval modals display exact line-by-line green additions and red deletions, ensuring zero ambiguity."),
        ("State Drift Detection", "Prior to rolling back, NEXUS checks if the target file was modified externally. If a drift conflict occurs, reversion is halted safely."),
        ("1-Click Rollback Execution", "Users can reverse mutations instantly from the Dashboard Audit Ledger or Ambient HUD, cleanly restoring directory trees."),
        ("Immutable Audit Ledger", "Every user request, LLM plan, tool execution, and token cost is permanently saved to an append-only database table.")
    ]
    for idx, (h, b) in enumerate(rev_points):
        p = tf10_r.paragraphs[0] if idx == 0 else tf10_r.add_paragraph()
        p.text = f"✔ {h}: "
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = GREEN_ACCENT
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = b
        run.font.bold = False
        run.font.color.rgb = TEXT_BODY

    # SLIDE 11: DEMERITS & SYSTEM LIMITATIONS
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Demerits & System Limitations", "Academic transparency: Key trade-offs, engineering constraints, and boundary conditions")

    demerits = [
        ("1. Multi-Step Execution Latency", "Reasoning & Coordination Overhead", [
            "Decomposing requests into DAGs and evaluating policy gates adds 1.5 - 3s latency compared to raw prompt chatbots.",
            "Complex tasks requiring multi-turn LLM reflections can increase end-to-end execution time for time-sensitive tasks."
        ]),
        ("2. External Side-Effect Irreversibility", "Network & API Boundaries", [
            "Filesystem changes are 100% reversible via snapshots, but external network actions (sent emails, dispatched webhooks) cannot be undone.",
            "Requires strict HIGH/CRITICAL policy classification to prevent accidental external dispatches."
        ]),
        ("3. OS Permission & Platform Dependency", "macOS TCC & Accessibility Gating", [
            "Native desktop control requires explicit macOS Accessibility (AXIsProcessTrusted) and Screen Recording TCC permissions.",
            "Cross-platform parity on Windows and Linux currently requires containerized or virtualized subprocess adapters."
        ]),
        ("4. Web DOM & UI Layout Drift", "Dynamic Web Application Fragility", [
            "Playwright browser automation can experience brittle selector matching if third-party web apps deploy sudden DOM changes.",
            "Mitigated by combining text content matching with resilient accessibility tree element resolution."
        ])
    ]

    for i, (title, sub, bullets) in enumerate(demerits):
        left_pos = Inches(0.8 + i * 2.98)
        card = add_card(s11, left_pos, Inches(1.8), Inches(2.8), Inches(5.1), title, border_color=RGBColor(254, 205, 211))
        tb = s11.shapes.add_textbox(left_pos + Inches(0.15), Inches(2.4), Inches(2.5), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        
        sp = tf.paragraphs[0]
        sp.text = sub
        sp.font.name = "Arial"
        sp.font.size = Pt(11)
        sp.font.bold = True
        sp.font.color.rgb = RED_ACCENT
        sp.space_after = Pt(8)

        for b in bullets:
            bp = tf.add_paragraph()
            bp.text = f"• {b}"
            bp.font.name = "Arial"
            bp.font.size = Pt(10)
            bp.font.color.rgb = TEXT_BODY
            bp.space_before = Pt(6)

    # SLIDE 12: FUTURE SCOPE & ROADMAP
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Future Scope & Engineering Roadmap", "Advancing from a developer platform toward ubiquitous agentic computing")

    roadmap_items = [
        ("Phase 16: Academic Evaluation Suite", "Formal Benchmark Ablation Experiments", [
            "Comprehensive benchmark harness testing across software engineering, sysadmin, and research workflows.",
            "Formal ablation studies: Baseline LLM vs. RAG vs. Memory-RAG vs. Full Agentic NEXUS.",
            "Quantitative measurement of Task Success Rate, Latency, Cost Efficiency, and Safety Violation Rate."
        ]),
        ("Multimodal Computer Vision Navigation", "Vision-in-the-Loop Desktop Control", [
            "Expanding beyond accessibility trees to real-time multimodal screen perception.",
            "Enables interaction with custom Canvas applications, legacy desktop GUIs, and interactive 3D software."
        ]),
        ("Full Cross-Platform OS Parity", "Linux & Windows Native Subsystems", [
            "Native Linux adapters via Wayland / X11 and D-Bus system bus integration.",
            "Windows support via UI Automation APIs and sandboxed PowerShell execution wrappers."
        ]),
        ("Cryptographically Verified Approvals", "Hardware-Gated Security Architecture", [
            "Touch ID / WebAuthn hardware security keys to sign CRITICAL-tier policy approvals.",
            "Zero-trust credential signing for production cloud deployments and irreversible operations."
        ])
    ]

    for i, (title, sub, bullets) in enumerate(roadmap_items):
        row = 0 if i < 2 else 1
        col = i % 2
        left_pos = Inches(0.8 + col * 5.95)
        top_pos = Inches(1.8 + row * 2.65)
        
        card = add_card(s12, left_pos, top_pos, Inches(5.6), Inches(2.4), title)
        tb = s12.shapes.add_textbox(left_pos + Inches(0.2), top_pos + Inches(0.55), Inches(5.2), Inches(1.7))
        tf = tb.text_frame
        tf.word_wrap = True
        
        sp = tf.paragraphs[0]
        sp.text = sub
        sp.font.name = "Arial"
        sp.font.size = Pt(11)
        sp.font.bold = True
        sp.font.color.rgb = ACCENT_BLUE
        sp.space_after = Pt(4)

        for b in bullets:
            bp = tf.add_paragraph()
            bp.text = f"• {b}"
            bp.font.name = "Arial"
            bp.font.size = Pt(11)
            bp.font.color.rgb = TEXT_BODY
            bp.space_before = Pt(3)

    # SLIDE 13: CONCLUSION
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, "Conclusion & Key Takeaways", "The synthesis of autonomous capability and uncompromised user safety")

    concl_banner = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(1.2))
    concl_banner.fill.solid()
    concl_banner.fill.fore_color.rgb = RGBColor(238, 242, 255)
    concl_banner.line.color.rgb = RGBColor(199, 210, 254)
    cbtb = s13.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.333), Inches(0.8))
    cbtf = cbtb.text_frame
    cbp = cbtf.paragraphs[0]
    cbp.text = "“Give AI the ability to act — without giving up control.”"
    cbp.alignment = PP_ALIGN.CENTER
    cbp.font.name = "Arial"
    cbp.font.size = Pt(22)
    cbp.font.bold = True
    cbp.font.color.rgb = NAVY

    cbp2 = cbtf.add_paragraph()
    cbp2.text = "Capability  +  Control  +  Observability  +  Reversibility  =  Trustworthy Agentic AI"
    cbp2.alignment = PP_ALIGN.CENTER
    cbp2.font.name = "Arial"
    cbp2.font.size = Pt(14)
    cbp2.font.bold = True
    cbp2.font.color.rgb = ACCENT_TEAL
    cbp2.space_before = Pt(4)

    concl_cards = [
        ("Proven Engineering Quality", [
            "15 completed development phases under a binding engineering constitution.",
            "194 automated unit and integration tests passing across all subsystems.",
            "Strict Pydantic v2 and TypeScript type safety with zero runtime mock bypasses."
        ]),
        ("True Agentic Operating System", [
            "Moves far beyond chatbots: natively operates terminals, browsers, and desktop software.",
            "Employs dynamic DAG replanning and self-correction to recover from runtime failures.",
            "Dual surfaces: Analytical Web Cockpit + Lightweight Ambient HUD."
        ]),
        ("Safe by Architectural Design", [
            "Air-gapped Policy Engine enforces 4-tier risk classification.",
            "Git-style unified diffs eliminate ambiguity during human approval.",
            "1-click snapshot restoration guarantees that user workspaces remain protected."
        ])
    ]

    for i, (title, bullets) in enumerate(concl_cards):
        left_pos = Inches(0.8 + i * 3.98)
        card = add_card(s13, left_pos, Inches(3.3), Inches(3.7), Inches(3.6), title)
        tb = s13.shapes.add_textbox(left_pos + Inches(0.15), Inches(3.9), Inches(3.4), Inches(2.8))
        tf = tb.text_frame
        tf.word_wrap = True
        
        for idx, b in enumerate(bullets):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = f"✔ {b}"
            p.font.name = "Arial"
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_BODY
            p.space_before = Pt(8)

    # SLIDE 14: Q&A AND THANK YOU
    s14 = prs.slides.add_slide(blank_layout)
    top_bar14 = s14.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
    top_bar14.fill.solid()
    top_bar14.fill.fore_color.rgb = ACCENT_BLUE
    top_bar14.line.fill.background()

    q_box = s14.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(3.5))
    qtf = q_box.text_frame
    qtf.word_wrap = True
    
    qp1 = qtf.paragraphs[0]
    qp1.text = "Thank You!"
    qp1.alignment = PP_ALIGN.CENTER
    qp1.font.name = "Arial"
    qp1.font.size = Pt(54)
    qp1.font.bold = True
    qp1.font.color.rgb = PRIMARY

    qp2 = qtf.add_paragraph()
    qp2.text = "Questions & Answers"
    qp2.alignment = PP_ALIGN.CENTER
    qp2.font.name = "Arial"
    qp2.font.size = Pt(28)
    qp2.font.bold = True
    qp2.font.color.rgb = ACCENT_BLUE
    qp2.space_before = Pt(8)

    qp3 = qtf.add_paragraph()
    qp3.text = "NEXUS: Agentic AI Operating System\nDepartment of Computer Science & Engineering • RVSCET, Jamshedpur"
    qp3.alignment = PP_ALIGN.CENTER
    qp3.font.name = "Arial"
    qp3.font.size = Pt(14)
    qp3.font.color.rgb = TEXT_MUTED
    qp3.space_before = Pt(14)

    add_card(s14, Inches(2.5), Inches(5.5), Inches(8.333), Inches(1.2), border_color=CARD_BORDER)
    pres_tb = s14.shapes.add_textbox(Inches(2.7), Inches(5.65), Inches(7.933), Inches(0.9))
    ptf_end = pres_tb.text_frame
    ptf_end.word_wrap = True
    
    pp_end = ptf_end.paragraphs[0]
    pp_end.text = "Group A5: Pushkar Raj (073) • Chandan Kr. Sah (074) • Prince Tiwari (075) • Shruti Kumari (076)"
    pp_end.alignment = PP_ALIGN.CENTER
    pp_end.font.name = "Arial"
    pp_end.font.size = Pt(13)
    pp_end.font.bold = True
    pp_end.font.color.rgb = PRIMARY

    pp_end2 = ptf_end.add_paragraph()
    pp_end2.text = "Guide: Sulekh Kumar (Assistant Professor, CSE)"
    pp_end2.alignment = PP_ALIGN.CENTER
    pp_end2.font.name = "Arial"
    pp_end2.font.size = Pt(12)
    pp_end2.font.color.rgb = TEXT_MUTED
    pp_end2.space_before = Pt(4)

    output_path = "/Users/apple/coding /nexus/NEXUS_Final_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation successfully saved to {output_path}")

if __name__ == "__main__":
    create_deck()
